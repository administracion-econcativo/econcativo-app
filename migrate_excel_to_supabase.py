"""
=============================================================================
SCRIPT DE MIGRACIÓN: EXCEL (operation_data.xlsx) -> SUPABASE (POSTGRESQL)
=============================================================================
Este script lee todas las hojas de operation_data.xlsx y las migra
de manera limpia y segura a la base de datos Supabase / PostgreSQL.

Uso:
  python migrate_excel_to_supabase.py          # Si DATABASE_URL está en .env
  python migrate_excel_to_supabase.py --dry-run # Para validar la lectura y mapeo sin conectar
"""

import os
import sys
import json
import re
from datetime import datetime, date
import openpyxl
from dotenv import load_dotenv

# Cargar variables de entorno desde .env
load_dotenv()

EXCEL_PATH = os.path.join(os.path.dirname(__file__), "operation_data.xlsx")
DATABASE_URL = os.environ.get("DATABASE_URL", "").strip()

def clean_val(v):
    if v is None:
        return None
    if isinstance(v, (datetime, date)):
        return v.strftime('%Y-%m-%d')
    s = str(v).strip()
    return s if s != "" else None

def clean_num(v, default=0.0):
    if v is None:
        return default
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace('$', '').replace(' ', '')
    if not s:
        return default
    try:
        if ',' in s and '.' in s:
            s = s.replace('.', '').replace(',', '.')
        elif ',' in s:
            s = s.replace(',', '.')
        return float(s)
    except:
        return default

def clean_int(v, default=None):
    if v is None:
        return default
    val = clean_num(v, default=None)
    if val is None:
        return default
    return int(val)

def normalize_key(k):
    s = str(k).lower().strip()
    s = s.replace('á', 'a').replace('é', 'e').replace('í', 'i').replace('ó', 'o').replace('ú', 'u')
    s = s.replace('ñ', 'n').replace('º', '').replace('°', '').replace('/', '_').replace('-', '_')
    s = re.sub(r'[^a-z0-9_]', '_', s)
    s = re.sub(r'_+', '_', s).strip('_')
    return s

def get_row_value(row_dict, *candidate_keys):
    for ck in candidate_keys:
        norm_ck = normalize_key(ck)
        for k, v in row_dict.items():
            if normalize_key(k) == norm_ck:
                return v
    return None


def read_sheet(wb, sheet_name):
    if sheet_name not in wb.sheetnames:
        return []
    ws = wb[sheet_name]
    raw_headers = [ws.cell(4, c).value for c in range(1, ws.max_column + 1)]
    headers = [str(h).strip() if h is not None else f"_col_{i}" for i, h in enumerate(raw_headers, 1)]
    
    rows = []
    for r in range(5, ws.max_row + 1):
        row_dict = {}
        has_data = False
        for c, h in enumerate(headers, 1):
            val = ws.cell(r, c).value
            if val is not None and str(val).strip() != "":
                has_data = True
            row_dict[h] = val
        if has_data:
            rows.append(row_dict)
    return rows

def parse_all_data(excel_path=EXCEL_PATH):
    print(f"[*] Abriendo archivo Excel: {excel_path}...")
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    parsed = {}

    # 1. USUARIOS
    usuarios = []
    for r in read_sheet(wb, 'USUARIOS'):
        u_id = clean_val(r.get('ID usuario'))
        if u_id:
            usuarios.append({
                'id_usuario': u_id,
                'nombre_completo': clean_val(r.get('Nombre completo')) or 'Usuario',
                'usuario': clean_val(r.get('Usuario')) or u_id,
                'password_hash': clean_val(r.get('Password Hash')) or '',
                'rol': clean_val(r.get('Rol')) or 'Administrador',
                'estado': clean_val(r.get('Estado')) or 'Activo',
                'ultimo_acceso': clean_val(r.get('Último acceso')) or clean_val(r.get('ltimo acceso'))
            })
    parsed['usuarios'] = usuarios

    # 2. CLIENTES
    clientes = []
    for r in read_sheet(wb, 'CLIENTES'):
        c_id = clean_val(r.get('ID cliente'))
        if c_id:
            clientes.append({
                'id_cliente': c_id,
                'razon_social': clean_val(r.get('Razón social / Nombre')) or clean_val(r.get('Razn social / Nombre')) or '',
                'cuit_cuil': clean_val(r.get('CUIT/CUIL')),
                'tipo_cliente': clean_val(r.get('Tipo de cliente')) or 'Empresa',
                'condicion_iva': clean_val(r.get('Condición IVA')) or clean_val(r.get('Condicin IVA')) or 'Responsable Inscripto',
                'telefono': clean_val(r.get('Teléfono')) or clean_val(r.get('Telfono')),
                'correo': clean_val(r.get('Correo')),
                'domicilio': clean_val(r.get('Domicilio')),
                'estado': clean_val(r.get('Estado')) or 'Activo',
                'observaciones': clean_val(r.get('Observaciones'))
            })
    parsed['clientes'] = clientes

    # 3. UNIDADES
    unidades = []
    for r in read_sheet(wb, 'UNIDADES'):
        u_id = clean_val(r.get('ID unidad'))
        if u_id:
            unidades.append({
                'id_unidad': u_id,
                'tipo': clean_val(r.get('Tipo')),
                'descripcion': clean_val(r.get('Descripción')) or clean_val(r.get('Descripcin')),
                'dominio_patente': clean_val(r.get('Dominio/Patente')) or u_id,
                'marca': clean_val(r.get('Marca')),
                'modelo': clean_val(r.get('Modelo')),
                'anio': clean_int(r.get('Año') or r.get('Ao'), default=None),
                'estado': clean_val(r.get('Estado')) or 'Activo',
                'seguro_vence': clean_val(r.get('Seguro vence')),
                'rto_itv_vence': clean_val(r.get('RTO/ITV vence')),
                'enlace_documentacion': clean_val(r.get('Enlace documentación')) or clean_val(r.get('Enlace documentacin')),
                'observaciones': clean_val(r.get('Observaciones'))
            })
    parsed['unidades'] = unidades

    # 4. EMPLEADOS
    empleados = []
    for r in read_sheet(wb, 'EMPLEADOS'):
        e_id = clean_val(r.get('ID empleado'))
        if e_id:
            empleados.append({
                'id_empleado': e_id,
                'apellido_nombre': clean_val(r.get('Apellido y nombre')) or '',
                'dni': clean_val(r.get('DNI')),
                'cuil': clean_val(r.get('CUIL')),
                'fecha_nacimiento': clean_val(r.get('Fecha de nacimiento')),
                'estado_civil': clean_val(r.get('Estado civil')),
                'domicilio_real_actual': clean_val(r.get('Domicilio real actual')),
                'barrio_zona': clean_val(r.get('Barrio / Zona')),
                'localidad': clean_val(r.get('Localidad')),
                'telefono': clean_val(r.get('Teléfono')) or clean_val(r.get('Telfono')),
                'email': clean_val(r.get('Email')),
                'contacto_emergencia': clean_val(r.get('Contacto de emergencia')),
                'vinculo_parentesco': clean_val(r.get('Vínculo / Parentesco')) or clean_val(r.get('Vnculo / Parentesco')),
                'domicilio_contacto': clean_val(r.get('Domicilio contacto')),
                'localidad_contacto': clean_val(r.get('Localidad contacto')),
                'fecha_ingreso': clean_val(r.get('Fecha de ingreso')),
                'puesto': clean_val(r.get('Puesto')),
                'modalidad_contratacion': clean_val(r.get('Modalidad de contratación')) or clean_val(r.get('Modalidad contratacin')),
                'sueldo_basico': clean_num(r.get('Sueldo básico') or r.get('Sueldo bsico')),
                'comision_pct': clean_num(r.get('Comisión %') or r.get('Comisin %')),
                'banco': clean_val(r.get('Banco')),
                'cbu': clean_val(r.get('CBU')),
                'alias_cbu': clean_val(r.get('Alias CBU')),
                'estado': clean_val(r.get('Estado')) or 'Activo',
                'observaciones': clean_val(r.get('Observaciones'))
            })
    parsed['empleados'] = empleados

    # 5. PROVEEDORES
    proveedores = []
    for r in read_sheet(wb, 'PROVEEDORES'):
        p_id = clean_val(r.get('ID proveedor'))
        if p_id:
            proveedores.append({
                'id_proveedor': p_id,
                'razon_social': clean_val(r.get('Razón social')) or clean_val(r.get('Razn social')) or '',
                'nombre_comercial': clean_val(r.get('Nombre comercial')),
                'cuit': clean_val(r.get('CUIT')),
                'rubro': clean_val(r.get('Rubro')),
                'telefono': clean_val(r.get('Teléfono')) or clean_val(r.get('Telfono')),
                'correo': clean_val(r.get('Correo')),
                'condicion_pago': clean_val(r.get('Condición de pago')) or clean_val(r.get('Condicin de pago')),
                'saldo_inicial': clean_num(r.get('Saldo inicial')),
                'fecha_saldo_inicial': clean_val(r.get('Fecha saldo inicial')),
                'estado': clean_val(r.get('Estado')) or 'Activo',
                'observaciones': clean_val(r.get('Observaciones')),
                'cbu_alias': clean_val(r.get('CBU / ALIAS'))
            })
    parsed['proveedores'] = proveedores

    # 6. ACTIVIDADES
    actividades = []
    for r in read_sheet(wb, 'ACTIVIDADES'):
        a_id = clean_val(r.get('ID actividad'))
        if a_id:
            actividades.append({
                'id_actividad': a_id,
                'actividad': clean_val(r.get('Actividad')) or '',
                'descripcion': clean_val(r.get('Descripción')) or clean_val(r.get('Descripcin')),
                'estado': clean_val(r.get('Estado')) or 'Activo'
            })
    parsed['actividades'] = actividades

    # 7. PRESUPUESTOS
    presupuestos = []
    for r in read_sheet(wb, 'PRESUPUESTOS'):
        p_id = clean_val(r.get('ID presupuesto'))
        if p_id:
            presupuestos.append({
                'id_presupuesto': p_id,
                'fecha': clean_val(r.get('Fecha')) or datetime.now().strftime('%Y-%m-%d'),
                'cliente': clean_val(r.get('Cliente')) or '',
                'cuit': clean_val(r.get('CUIT')),
                'tipo_servicio': clean_val(r.get('Tipo servicio')),
                'detalle_concepto': clean_val(r.get('Detalle/Concepto')),
                'total': clean_num(r.get('Total')),
                'iva_pct': clean_num(r.get('IVA %')),
                'validez': clean_val(r.get('Validez')),
                'forma_pago': clean_val(r.get('Forma pago')),
                'estado': clean_val(r.get('Estado')) or 'Pendiente',
                'id_viaje_asociado': clean_val(r.get('ID viaje asociado')),
                'observaciones': clean_val(r.get('Observaciones'))
            })
    parsed['presupuestos'] = presupuestos

    # 8. VIAJES
    viajes = []
    for r in read_sheet(wb, 'VIAJES'):
        v_id = clean_val(r.get('ID viaje'))
        if v_id:
            viajes.append({
                'id_viaje': v_id,
                'fecha_salida': clean_val(r.get('Fecha salida')) or datetime.now().strftime('%Y-%m-%d'),
                'fecha_regreso': clean_val(r.get('Fecha regreso')),
                'chofer': clean_val(r.get('Chofer')),
                'unidad': clean_val(r.get('Unidad')),
                'origen': clean_val(r.get('Origen')),
                'destino': clean_val(r.get('Destino')),
                'cliente': clean_val(r.get('Cliente')) or '',
                'actividad': clean_val(r.get('Actividad')),
                'carga': clean_val(r.get('Carga')),
                'nro_remito': clean_val(r.get('Nº remito')) or clean_val(r.get('N remito')),
                'cpe': clean_val(r.get('CPE')),
                'km_inicial': clean_num(r.get('Km inicial'), default=None),
                'km_final': clean_num(r.get('Km final'), default=None),
                'km_recorridos': clean_num(r.get('Km recorridos'), default=None),
                'toneladas': clean_num(r.get('Toneladas'), default=None),
                'combustible_litros': clean_num(r.get('Combustible litros'), default=None),
                'gastos_viaje': clean_num(r.get('Gastos viaje')),
                'adelanto': clean_num(r.get('Adelanto')),
                'valor_servicio': clean_num(r.get('Valor servicio')),
                'estado_facturacion': clean_val(r.get('Estado facturación')) or clean_val(r.get('Estado facturacin')) or 'No facturado',
                'id_ingreso_factura': clean_val(r.get('ID ingreso/factura')),
                'observaciones': clean_val(r.get('Observaciones'))
            })
    parsed['viajes'] = viajes

    # 9. INGRESOS
    ingresos = []
    for r in read_sheet(wb, 'INGRESOS'):
        i_id = clean_val(r.get('ID ingreso'))
        if i_id:
            ingresos.append({
                'id_ingreso': i_id,
                'fecha_emision': clean_val(r.get('Fecha emisión')) or clean_val(r.get('Fecha emisin')) or datetime.now().strftime('%Y-%m-%d'),
                'tipo_comprobante': clean_val(r.get('Tipo comprobante')),
                'punto_venta': clean_val(r.get('Punto venta')),
                'nro_factura_arca': clean_val(r.get('Nº factura ARCA')) or clean_val(r.get('N factura ARCA')),
                'cliente': clean_val(r.get('Cliente')) or '',
                'cuit_cuil': clean_val(r.get('CUIT/CUIL')),
                'actividad': clean_val(r.get('Actividad')),
                'chofer': clean_val(r.get('Chofer')),
                'unidad': clean_val(r.get('Unidad')),
                'km': clean_num(r.get('Km'), default=None),
                'toneladas': clean_num(r.get('Toneladas'), default=None),
                'neto': clean_num(r.get('Neto')),
                'iva_pct': clean_num(r.get('IVA %')),
                'iva': clean_num(r.get('IVA')),
                'total': clean_num(r.get('Total')),
                'estado_cobro': clean_val(r.get('Estado cobro')) or 'PENDIENTE',
                'importe_cobrado': clean_num(r.get('Importe cobrado')),
                'fecha_cobro': clean_val(r.get('Fecha cobro')),
                'medio_cobro': clean_val(r.get('Medio cobro')),
                'saldo': clean_num(r.get('Saldo')),
                'id_acuerdo': clean_val(r.get('ID acuerdo')),
                'modalidad_aplicada': clean_val(r.get('Modalidad aplicada')),
                'tarifa_aplicada': clean_num(r.get('Tarifa aplicada'), default=None),
                'remuneracion': clean_num(r.get('Remuneración') or r.get('Remuneracin'), default=None),
                'estado_remuneracion': clean_val(r.get('Estado remuneración') or r.get('Estado remuneracin')),
                'enlace_factura': clean_val(r.get('Enlace factura')),
                'observaciones': clean_val(r.get('Observaciones'))
            })
    parsed['ingresos'] = ingresos

    # 10. EGRESOS
    egresos = []
    for r in read_sheet(wb, 'EGRESOS'):
        e_id = clean_val(r.get('ID egreso'))
        if e_id:
            egresos.append({
                'id_egreso': e_id,
                'fecha_comprobante': clean_val(r.get('Fecha comprobante')) or datetime.now().strftime('%Y-%m-%d'),
                'fecha_vencimiento': clean_val(r.get('Fecha vencimiento')),
                'categoria': clean_val(r.get('Categoría')) or clean_val(r.get('Categora')),
                'subcategoria': clean_val(r.get('Subcategoría')) or clean_val(r.get('Subcategora')),
                'proveedor': clean_val(r.get('Proveedor')),
                'cuit': clean_val(r.get('CUIT')),
                'tipo_comprobante': clean_val(r.get('Tipo comprobante')),
                'nro_comprobante': clean_val(r.get('Nº comprobante')) or clean_val(r.get('N comprobante')),
                'unidad': clean_val(r.get('Unidad')),
                'empleado': clean_val(r.get('Empleado')),
                'descripcion': clean_val(r.get('Descripción')) or clean_val(r.get('Descripcin')),
                'cantidad_litros': clean_num(r.get('Cantidad/Litros'), default=None),
                'neto': clean_num(r.get('Neto')),
                'iva': clean_num(r.get('IVA')),
                'total': clean_num(r.get('Total')),
                'estado_pago': clean_val(r.get('Estado pago')) or 'Pendiente',
                'importe_pagado': clean_num(r.get('Importe pagado')),
                'fecha_pago': clean_val(r.get('Fecha pago')),
                'medio_pago': clean_val(r.get('Medio pago')),
                'saldo': clean_num(r.get('Saldo')),
                'enlace_comprobante': clean_val(r.get('Enlace comprobante')),
                'id_liquidacion': clean_val(r.get('ID liquidación')) or clean_val(r.get('ID liquidacin')),
                'observaciones': clean_val(r.get('Observaciones'))
            })
    parsed['egresos'] = egresos

    # 11. CHEQUES
    cheques = []
    for r in read_sheet(wb, 'CHEQUES'):
        ch_id = clean_val(r.get('ID cheque'))
        if ch_id:
            cheques.append({
                'id_cheque': ch_id,
                'fecha_ingreso': clean_val(r.get('Fecha ingreso')) or datetime.now().strftime('%Y-%m-%d'),
                'tipo': clean_val(r.get('Tipo')),
                'monto': clean_num(r.get('Monto')),
                'cliente_emisor': clean_val(r.get('Cliente / Emisor')),
                'cuit_emisor': clean_val(r.get('CUIT emisor')),
                'banco': clean_val(r.get('Banco')),
                'nro_cheque': clean_val(r.get('Nº Cheque')) or clean_val(r.get('N Cheque')),
                'fecha_cobro': clean_val(r.get('Fecha cobro')),
                'endosado_tenedor': clean_val(r.get('Endosado / Tenedor')),
                'estado': clean_val(r.get('Estado')) or 'En Cartera',
                'destino_usado_en': clean_val(r.get('Destino / Usado en')),
                'fecha_uso': clean_val(r.get('Fecha uso')),
                'id_egreso': clean_val(r.get('ID egreso')),
                'id_ingreso_origen': clean_val(r.get('ID ingreso origen')),
                'observaciones': clean_val(r.get('Observaciones'))
            })
    parsed['cheques'] = cheques

    # 12. ORDENES_COMPRA
    ordenes = []
    for r in read_sheet(wb, 'ORDENES_COMPRA'):
        o_id = clean_val(r.get('ID orden'))
        if o_id:
            ordenes.append({
                'id_orden': o_id,
                'fecha': clean_val(r.get('Fecha')) or datetime.now().strftime('%Y-%m-%d'),
                'proveedor': clean_val(r.get('Proveedor')) or '',
                'cuit_proveedor': clean_val(r.get('CUIT proveedor')),
                'tipo_insumo': clean_val(r.get('Tipo insumo')),
                'unidad': clean_val(r.get('Unidad')),
                'detalle': clean_val(r.get('Detalle')),
                'neto': clean_num(r.get('Neto')),
                'iva_pct': clean_num(r.get('IVA %')),
                'total': clean_num(r.get('Total')),
                'forma_pago': clean_val(r.get('Forma pago')),
                'estado': clean_val(r.get('Estado')) or 'Pendiente',
                'fecha_pago': clean_val(r.get('Fecha pago')),
                'id_egreso': clean_val(r.get('ID egreso')),
                'observaciones': clean_val(r.get('Observaciones')),
                'cantidad_litros': clean_num(r.get('Cantidad/Litros'), default=None),
                'id_cheque': clean_val(r.get('ID cheque')),
                'nro_cheque': clean_val(r.get('Nº cheque')) or clean_val(r.get('N cheque'))
            })
    parsed['ordenes_compra'] = ordenes

    # 13. LIQUIDACIONES
    liquidaciones = []
    for r in read_sheet(wb, 'LIQUIDACIONES'):
        l_id = clean_val(r.get('ID liquidación')) or clean_val(r.get('ID liquidacin'))
        if l_id:
            desglose = r.get('Desglose JSON')
            if isinstance(desglose, str) and (desglose.startswith('{') or desglose.startswith('[')):
                try:
                    desglose = json.loads(desglose)
                except:
                    desglose = None
            else:
                desglose = None
            liquidaciones.append({
                'id_liquidacion': l_id,
                'fecha_confirmacion': clean_val(r.get('Fecha confirmación')) or clean_val(r.get('Fecha confirmacin')),
                'periodo_quincena': clean_val(r.get('Período quincena')) or clean_val(r.get('Perodo quincena')),
                'id_empleado': clean_val(r.get('ID empleado')),
                'empleado': clean_val(r.get('Empleado')),
                'puesto': clean_val(r.get('Puesto')),
                'dni': clean_val(r.get('DNI')),
                'cuil': clean_val(r.get('CUIL')),
                'entidad_bancaria': clean_val(r.get('Entidad bancaria')),
                'cbu': clean_val(r.get('CBU')),
                'alias_cbu': clean_val(r.get('Alias CBU')),
                'sueldo_fijo_base': clean_num(r.get('Sueldo fijo base')),
                'comision_pct': clean_num(r.get('Comisión pct') or r.get('Comisin pct')),
                'cant_viajes': clean_int(r.get('Cant viajes')),
                'total_comisiones': clean_num(r.get('Total comisiones')),
                'sueldo_final': clean_num(r.get('Sueldo final')),
                'estado': clean_val(r.get('Estado')) or 'PENDIENTE',
                'desglose_json': json.dumps(desglose) if desglose is not None else None
            })
    parsed['liquidaciones'] = liquidaciones

    # 14. CUENTAS_CORRIENTES
    cc = []
    for r in read_sheet(wb, 'CUENTAS_CORRIENTES'):
        m_id = clean_val(r.get('ID movimiento'))
        if m_id:
            cc.append({
                'id_movimiento': m_id,
                'fecha': clean_val(r.get('Fecha')) or datetime.now().strftime('%Y-%m-%d'),
                'fecha_vencimiento': clean_val(r.get('Fecha vencimiento')),
                'cliente': clean_val(r.get('Cliente')) or '',
                'cuit': clean_val(r.get('CUIT')),
                'tipo_comprobante': clean_val(r.get('Tipo comprobante')),
                'referencia_id': clean_val(r.get('Referencia ID')),
                'concepto': clean_val(r.get('Concepto')),
                'debe': clean_num(r.get('Debe')),
                'haber': clean_num(r.get('Haber')),
                'observaciones': clean_val(r.get('Observaciones'))
            })
    parsed['cuentas_corrientes'] = cc

    # 15. NOVEDADES_PERSONAL
    nov = []
    for r in read_sheet(wb, 'NOVEDADES_PERSONAL'):
        n_id = clean_val(r.get('ID novedad'))
        if n_id:
            nov.append({
                'id_novedad': n_id,
                'fecha': clean_val(r.get('Fecha')) or datetime.now().strftime('%Y-%m-%d'),
                'id_empleado': clean_val(r.get('ID empleado')),
                'empleado': clean_val(r.get('Empleado')),
                'tipo': clean_val(r.get('Tipo')),
                'concepto': clean_val(r.get('Concepto')),
                'monto': clean_num(r.get('Monto')),
                'medio_pago': clean_val(r.get('Medio pago')),
                'id_cheque': clean_val(r.get('ID cheque')),
                'nro_cheque': clean_val(r.get('Nº cheque')) or clean_val(r.get('N cheque')),
                'id_liquidacion': clean_val(r.get('ID liquidacion')),
                'estado': clean_val(r.get('Estado')) or 'Pendiente',
                'observaciones': clean_val(r.get('Observaciones'))
            })
    parsed['novedades_personal'] = nov

    # 16. TESORERIA CUENTAS
    tesoreria = []
    for r in read_sheet(wb, 'TESORERIA'):
        c_id = clean_val(r.get('ID cuenta'))
        if c_id:
            tesoreria.append({
                'id_cuenta': c_id,
                'nombre_cuenta': clean_val(r.get('Nombre cuenta')) or c_id,
                'tipo': clean_val(r.get('Tipo')),
                'nro_cuenta_cbu_alias': clean_val(r.get('Nº Cuenta / CBU / Alias')) or clean_val(r.get('N Cuenta / CBU / Alias')),
                'saldo_inicial': clean_num(r.get('Saldo inicial')),
                'saldo_real_cotejado': clean_num(r.get('Saldo real cotejado'), default=None),
                'fecha_ultimo_cotejo': clean_val(r.get('Fecha ultimo cotejo')),
                'observaciones': clean_val(r.get('Observaciones'))
            })
    parsed['tesoreria_cuentas'] = tesoreria

    # 17. CONFIGURACION
    config_items = []
    if 'CONFIGURACION' in wb.sheetnames:
        ws_cfg = wb['CONFIGURACION']
        col_headers = [ws_cfg.cell(4, c).value for c in range(1, ws_cfg.max_column + 1)]
        for c_idx, h in enumerate(col_headers, 1):
            if not h:
                continue
            cat_name = normalize_key(h)
            for r in range(5, ws_cfg.max_row + 1):
                val = ws_cfg.cell(r, c_idx).value
                if val is not None and str(val).strip() != "":
                    config_items.append({
                        'categoria': cat_name,
                        'clave': str(val).strip(),
                        'valor': str(val).strip(),
                        'orden': r - 4
                    })
    parsed['configuracion'] = config_items

    return parsed

def migrate_to_database(parsed_data, db_url):
    import psycopg2
    from psycopg2.extras import execute_batch

    print(f"\n[*] Conectando a Supabase PostgreSQL...")
    conn = psycopg2.connect(db_url)
    conn.autocommit = False
    cur = conn.cursor()

    try:
        # Ejecutar schema.sql primero si existe
        schema_path = os.path.join(os.path.dirname(__file__), "schema.sql")
        if os.path.exists(schema_path):
            print("[*] Verificando / creando tablas con schema.sql...")
            with open(schema_path, encoding='utf-8') as f:
                cur.execute(f.read())
            conn.commit()

        # Ajustar columnas numéricas en tablas existentes para evitar desbordes
        try:
            cur.execute("""
                DO $$
                BEGIN
                    IF EXISTS (SELECT 1 FROM information_schema.columns WHERE table_name = 'liquidaciones' AND column_name = 'comision_pct') THEN
                        ALTER TABLE liquidaciones ALTER COLUMN comision_pct TYPE NUMERIC(24, 2);
                    END IF;
                    IF EXISTS (SELECT 1 FROM information_schema.columns WHERE table_name = 'empleados' AND column_name = 'comision_pct') THEN
                        ALTER TABLE empleados ALTER COLUMN comision_pct TYPE NUMERIC(24, 2);
                    END IF;
                END $$;
            """)
            conn.commit()
        except Exception:
            conn.rollback()

        # Inserción por cada tabla usando ON CONFLICT DO UPDATE
        for table, rows in parsed_data.items():
            if not rows:
                print(f"[-] {table}: 0 registros para migrar.")
                continue

            first_row = rows[0]
            columns = list(first_row.keys())
            cols_str = ", ".join(columns)
            placeholders = ", ".join([f"%({c})s" for c in columns])
            
            # Primary key mapping
            pk_map = {
                'usuarios': 'id_usuario',
                'clientes': 'id_cliente',
                'unidades': 'id_unidad',
                'empleados': 'id_empleado',
                'proveedores': 'id_proveedor',
                'actividades': 'id_actividad',
                'presupuestos': 'id_presupuesto',
                'viajes': 'id_viaje',
                'ingresos': 'id_ingreso',
                'egresos': 'id_egreso',
                'cheques': 'id_cheque',
                'ordenes_compra': 'id_orden',
                'liquidaciones': 'id_liquidacion',
                'cuentas_corrientes': 'id_movimiento',
                'novedades_personal': 'id_novedad',
                'tesoreria_cuentas': 'id_cuenta',
                'configuracion': None
            }

            pk = pk_map.get(table)
            if pk:
                update_cols = [f"{c} = EXCLUDED.{c}" for c in columns if c != pk]
                update_str = ", ".join(update_cols)
                query = f"""
                    INSERT INTO {table} ({cols_str})
                    VALUES ({placeholders})
                    ON CONFLICT ({pk}) DO UPDATE SET {update_str};
                """
            else:
                query = f"""
                    INSERT INTO {table} ({cols_str})
                    VALUES ({placeholders})
                    ON CONFLICT DO NOTHING;
                """

            execute_batch(cur, query, rows)
            print(f"[+] {table}: {len(rows)} registros migrados con éxito.")

        conn.commit()
        print("\n[OK] MIGRACION COMPLETADA EXITOSAMENTE EN SUPABASE!")

    except Exception as e:
        conn.rollback()
        print(f"\n[ERROR] Error durante la migracion: {e}")
        raise
    finally:
        cur.close()
        conn.close()

if __name__ == '__main__':
    dry_run = '--dry-run' in sys.argv
    print("=" * 70)
    print("INICIANDO PROCESO DE MIGRACIÓN: EXCEL -> SUPABASE")
    print("=" * 70)

    data = parse_all_data()

    print("\n" + "=" * 70)
    print("RESUMEN DE REGISTROS LEÍDOS DESDE EXCEL:")
    print("=" * 70)
    total_records = 0
    for table_name, rows in data.items():
        count = len(rows)
        total_records += count
        print(f"  • {table_name:<25}: {count:>4} filas")
    print("-" * 70)
    print(f"  TOTAL GENERAL             : {total_records:>4} registros")
    print("=" * 70)

    if dry_run or not DATABASE_URL:
        if not DATABASE_URL:
            print("\n[i] AVISO: DATABASE_URL no está configurado en el archivo .env.")
            print("[i] Se ejecutó en modo DRY-RUN (Validación de lectura y tipos).")
            print("[i] Para migrar directamente a Supabase:")
            print("    1. Agregá tu DATABASE_URL en el archivo .env")
            print("    2. Volvé a ejecutar: python migrate_excel_to_supabase.py")
        else:
            print("\n[i] Modo DRY-RUN ejecutado con éxito. No se escribió en la base de datos.")
    else:
        migrate_to_database(data, DATABASE_URL)
